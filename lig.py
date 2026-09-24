import streamlit as st
import json, os, sqlite3
from contextlib import contextmanager

st.set_page_config(page_title="Lig Yöneticisi", page_icon="⚽", layout="wide")

# ── DB yolu ───────────────────────────────────────────────────────────────────
_local = os.path.join(os.path.dirname(os.path.abspath(__file__)), "lig.db")
DB_FILE = _local if os.access(os.path.dirname(_local), os.W_OK) else "/tmp/lig.db"
_json_local = os.path.join(os.path.dirname(os.path.abspath(__file__)), "lig_data.json")
JSON_FILE = _json_local if os.path.exists(_json_local) else "/tmp/lig_data.json"

# ── SQLite ────────────────────────────────────────────────────────────────────
@contextmanager
def get_db():
    conn = sqlite3.connect(DB_FILE)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA journal_mode=WAL")
    try:
        yield conn; conn.commit()
    except Exception:
        conn.rollback(); raise
    finally:
        conn.close()

def init_db():
    with get_db() as conn:
        conn.executescript("""
            CREATE TABLE IF NOT EXISTS teams (
                id   INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL UNIQUE
            );
            CREATE TABLE IF NOT EXISTS fixtures (
                id     INTEGER PRIMARY KEY AUTOINCREMENT,
                week   INTEGER NOT NULL,
                devre  INTEGER NOT NULL DEFAULT 1,
                home   TEXT NOT NULL,
                away   TEXT NOT NULL,
                hg     INTEGER,
                ag     INTEGER,
                played INTEGER NOT NULL DEFAULT 0
            );
        """)

def migrate_from_json():
    if not os.path.exists(JSON_FILE): return
    try:
        with open(JSON_FILE, "r", encoding="utf-8") as f:
            d = json.load(f)
        if not d.get("teams") and not d.get("fixtures"): return
        with get_db() as conn:
            if conn.execute("SELECT COUNT(*) FROM teams").fetchone()[0] > 0: return
            for t in d.get("teams", []):
                conn.execute("INSERT OR IGNORE INTO teams (name) VALUES (?)", (t,))
            for m in d.get("fixtures", []):
                conn.execute(
                    "INSERT OR IGNORE INTO fixtures (id,week,devre,home,away,hg,ag,played) VALUES (?,?,?,?,?,?,?,?)",
                    (m["id"], m["week"], m.get("devre",1), m["home"], m["away"],
                     m.get("hg"), m.get("ag"), 1 if m.get("played") else 0)
                )
        done = JSON_FILE.replace(".json", "_migrated.json")
        os.rename(JSON_FILE, done)
        st.toast("✅ JSON → SQLite taşındı!", icon="🎉")
    except Exception as e:
        st.warning(f"Migration hatası: {e}")

def db_get_teams():
    with get_db() as conn:
        return [r["name"] for r in conn.execute("SELECT name FROM teams ORDER BY id").fetchall()]

def db_get_fixtures():
    with get_db() as conn:
        rows = conn.execute("SELECT * FROM fixtures ORDER BY week, id").fetchall()
    return [dict(r) | {"played": bool(r["played"])} for r in rows]

def db_save_team(name):
    with get_db() as conn:
        conn.execute("INSERT INTO teams (name) VALUES (?)", (name,))

def db_rename_team(old, new):
    with get_db() as conn:
        conn.execute("UPDATE teams SET name=? WHERE name=?", (new, old))
        conn.execute("UPDATE fixtures SET home=? WHERE home=?", (new, old))
        conn.execute("UPDATE fixtures SET away=? WHERE away=?", (new, old))

def db_delete_team(name):
    with get_db() as conn:
        conn.execute("DELETE FROM teams WHERE name=?", (name,))
        conn.execute("DELETE FROM fixtures WHERE home=? OR away=?", (name, name))

def db_add_fixture(week, devre, home, away, hg=None, ag=None, played=False):
    with get_db() as conn:
        conn.execute(
            "INSERT INTO fixtures (week,devre,home,away,hg,ag,played) VALUES (?,?,?,?,?,?,?)",
            (week, devre, home, away, hg, ag, 1 if played else 0)
        )

def db_bulk_add_fixtures(rows):
    with get_db() as conn:
        conn.executemany(
            "INSERT INTO fixtures (week,devre,home,away,hg,ag,played) VALUES (?,?,?,?,NULL,NULL,0)",
            rows
        )

def db_update_score(fid, hg, ag, played):
    with get_db() as conn:
        conn.execute("UPDATE fixtures SET hg=?,ag=?,played=? WHERE id=?",
                     (hg, ag, 1 if played else 0, fid))

def db_update_fixture_teams(fid, home, away):
    with get_db() as conn:
        conn.execute("UPDATE fixtures SET home=?,away=? WHERE id=?", (home, away, fid))

def db_delete_fixture(fid):
    with get_db() as conn:
        conn.execute("DELETE FROM fixtures WHERE id=?", (fid,))

def db_reset_scores():
    with get_db() as conn:
        conn.execute("UPDATE fixtures SET hg=NULL,ag=NULL,played=0")

def db_reset_all():
    with get_db() as conn:
        conn.execute("DELETE FROM fixtures")
        conn.execute("DELETE FROM teams")

def db_clear_fixtures():
    with get_db() as conn:
        conn.execute("DELETE FROM fixtures")

# ── Uyumluluk: data dict yerine doğrudan listeler ─────────────────────────────
def make_data(teams, fixtures):
    return {"n": len(teams), "teams": teams, "fixtures": fixtures}

# ── Yardımcı ─────────────────────────────────────────────────────────────────
def get_devre(week, fixtures):
    tum_haftalar = sorted(set(m["week"] for m in fixtures) | {week})
    yari = len(tum_haftalar) // 2
    return 1 if week <= yari else 2

def next_id(fixtures):
    return max((m["id"] for m in fixtures), default=0) + 1

# ── İstatistik ────────────────────────────────────────────────────────────────
def compute_table(data):
    teams = data["teams"]
    if not teams:
        return []
    stats = {t: {"O":0,"G":0,"B":0,"M":0,"AG":0,"YG":0,
                 "IC_O":0,"IC_G":0,"IC_B":0,"IC_M":0,"IC_AG":0,"IC_YG":0,
                 "DIS_O":0,"DIS_G":0,"DIS_B":0,"DIS_M":0,"DIS_AG":0,"DIS_YG":0,
                 "form":[]} for t in teams}
    for m in data["fixtures"]:
        if not m["played"]:
            continue
        h, a = m["home"], m["away"]
        hg, ag = m["hg"], m["ag"]
        if h not in stats or a not in stats:
            continue
        stats[h]["O"] += 1; stats[a]["O"] += 1
        stats[h]["AG"] += hg; stats[h]["YG"] += ag
        stats[a]["AG"] += ag; stats[a]["YG"] += hg
        # İç saha
        stats[h]["IC_O"]  += 1
        stats[h]["IC_AG"] += hg
        stats[h]["IC_YG"] += ag
        # Dış saha
        stats[a]["DIS_O"]  += 1
        stats[a]["DIS_AG"] += ag
        stats[a]["DIS_YG"] += hg
        if hg > ag:
            stats[h]["G"] += 1; stats[a]["M"] += 1
            stats[h]["IC_G"]  += 1
            stats[a]["DIS_M"] += 1
            stats[h]["form"].append(("G", f"{hg}-{ag}", a))
            stats[a]["form"].append(("M", f"{ag}-{hg}", h))
        elif hg < ag:
            stats[a]["G"] += 1; stats[h]["M"] += 1
            stats[h]["IC_M"]  += 1
            stats[a]["DIS_G"] += 1
            stats[h]["form"].append(("M", f"{hg}-{ag}", a))
            stats[a]["form"].append(("G", f"{ag}-{hg}", h))
        else:
            stats[h]["B"] += 1; stats[a]["B"] += 1
            stats[h]["IC_B"]  += 1
            stats[a]["DIS_B"] += 1
            stats[h]["form"].append(("B", f"{hg}-{ag}", a))
            stats[a]["form"].append(("B", f"{hg}-{ag}", h))
    rows = []
    for t in teams:
        s = stats[t]
        puan = s["G"] * 3 + s["B"]
        av = s["AG"] - s["YG"]
        rows.append({
            "Takım": t, "O": s["O"], "G": s["G"], "B": s["B"], "M": s["M"],
            "AG": s["AG"], "YG": s["YG"], "AV": av, "Puan": puan,
            "İç O": s["IC_O"], "İç G": s["IC_G"], "İç B": s["IC_B"], "İç M": s["IC_M"],
            "İç AG": s["IC_AG"], "İç YG": s["IC_YG"],
            "Dış O": s["DIS_O"], "Dış G": s["DIS_G"], "Dış B": s["DIS_B"], "Dış M": s["DIS_M"],
            "Dış AG": s["DIS_AG"], "Dış YG": s["DIS_YG"],
            "form": s["form"][-5:]
        })
    rows.sort(key=lambda x: (-x["Puan"], -x["AV"], -x["AG"]))
    for i, r in enumerate(rows):
        r["Sıra"] = i + 1
    return rows

def export_excel(teams, fixtures):
    """3 sayfalı Excel dosyası üret: Puan Tablosu, Fikstür, İstatistikler"""
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    # ── Renkler ──────────────────────────────────────────────────────────
    C_HEADER  = "1F4E79"
    C_SUBHDR  = "2E75B6"
    C_GREEN   = "C6EFCE"; C_GREEN_F = "276221"
    C_RED     = "FFC7CE"; C_RED_F   = "9C0006"
    C_YELLOW  = "FFF2CC"
    C_ODD     = "DEEAF1"
    C_GRP     = "4472C4"  # grup ayırıcı

    def hfont(sz=10, bold=True):
        return Font(name="Arial", bold=bold, color="FFFFFF", size=sz)
    def cfont(sz=10, bold=False, color="000000"):
        return Font(name="Arial", size=sz, bold=bold, color=color)
    def fill(c):
        return PatternFill("solid", fgColor=c)
    def ctr():
        return Alignment(horizontal="center", vertical="center")
    def lft():
        return Alignment(horizontal="left", vertical="center")
    def border():
        s = Side(style="thin", color="CCCCCC")
        return Border(left=s, right=s, top=s, bottom=s)
    def grp_border():
        thin = Side(style="thin", color="CCCCCC")
        med  = Side(style="medium", color=C_GRP)
        return Border(left=med, right=thin, top=thin, bottom=thin)

    table = compute_table(make_data(teams, fixtures))

    # ══════════════════════════════════════════════════════════════════════
    # SAYFA 1 — PUAN TABLOSU
    # ══════════════════════════════════════════════════════════════════════
    ws1 = wb.active
    ws1.title = "Puan Tablosu"
    ws1.sheet_view.showGridLines = False
    ws1.freeze_panes = "A3"

    # Başlık
    ws1.merge_cells("A1:X1")
    c = ws1["A1"]
    c.value = "⚽  LİG PUAN TABLOSU"
    c.font = Font(name="Arial", bold=True, size=14, color="FFFFFF")
    c.fill = fill(C_HEADER)
    c.alignment = ctr()
    ws1.row_dimensions[1].height = 28

    # Sütun başlıkları — 3 grup
    grp1 = ["Sıra","Takım","Oyn","Gal","Ber","Mağ","A Gol","Y Gol","Aver","Puan"]
    grp2 = ["İç Oyn","İç Gal","İç Ber","İç Mağ","İç A Gol","İç Y Gol"]
    grp3 = ["Dış Oyn","Dış Gal","Dış Ber","Dış Mağ","Dış A Gol","Dış Y Gol"]
    grp4 = ["Son 5"]
    all_cols = grp1 + grp2 + grp3 + grp4
    grp2_start = len(grp1) + 1
    grp3_start = grp2_start + len(grp2)
    grp4_start = grp3_start + len(grp3)

    widths = [6,22,6,6,6,6,7,7,7,7, 8,7,7,7,8,8, 8,7,7,7,8,8, 14]
    for ci, (h, w) in enumerate(zip(all_cols, widths), 1):
        c = ws1.cell(2, ci, h)
        is_grp_start = ci in (grp2_start, grp3_start, grp4_start)
        c.font  = Font(name="Arial", bold=True, color="FFFFFF", size=9)
        c.fill  = fill(C_GRP if is_grp_start else C_SUBHDR)
        c.alignment = ctr()
        c.border = border()
        ws1.column_dimensions[get_column_letter(ci)].width = w
    ws1.row_dimensions[2].height = 20

    # Veri satırları
    key_map = ["Sıra","Takım","O","G","B","M","AG","YG","AV","Puan",
               "İç O","İç G","İç B","İç M","İç AG","İç YG",
               "Dış O","Dış G","Dış B","Dış M","Dış AG","Dış YG"]

    for ri, row in enumerate(table):
        r = ri + 3
        bg = C_ODD if ri % 2 == 0 else "FFFFFF"
        s = row["Sıra"]

        for ci, key in enumerate(key_map, 1):
            c = ws1.cell(r, ci, row[key])
            c.font      = cfont(bold=(key == "Puan"))
            c.alignment = lft() if key == "Takım" else ctr()
            c.border    = grp_border() if ci in (grp2_start, grp3_start) else border()
            # Zona rengi
            if s <= 3:
                c.fill = fill(C_GREEN)
                c.font = Font(name="Arial", size=10, color=C_GREEN_F, bold=(key=="Puan"))
            elif s >= len(table) - 2:
                c.fill = fill(C_RED)
                c.font = Font(name="Arial", size=10, color=C_RED_F, bold=(key=="Puan"))
            elif 4 <= s <= 6:
                c.fill = fill(C_YELLOW)
            else:
                c.fill = fill(bg)

        # Son 5 form — metin olarak
        form5 = row.get("form", [])
        form_str = " ".join(
            (f[0] if isinstance(f, tuple) else f) for f in form5
        )
        c = ws1.cell(r, 23, form_str)
        c.alignment = ctr()
        c.border = grp_border()
        c.fill = fill(C_GREEN if s <= 3 else (C_RED if s >= len(table)-2 else (C_YELLOW if 4<=s<=6 else bg)))
        ws1.row_dimensions[r].height = 16

    ws1.sheet_properties.tabColor = C_HEADER

    # ══════════════════════════════════════════════════════════════════════
    # SAYFA 2 — FİKSTÜR
    # ══════════════════════════════════════════════════════════════════════
    ws2 = wb.create_sheet("Fikstür")
    ws2.sheet_view.showGridLines = False
    ws2.freeze_panes = "A3"

    ws2.merge_cells("A1:H1")
    c = ws2["A1"]
    c.value = "📅  FİKSTÜR"
    c.font = Font(name="Arial", bold=True, size=14, color="FFFFFF")
    c.fill = fill(C_HEADER); c.alignment = ctr()
    ws2.row_dimensions[1].height = 28

    fk_hdrs = ["No","Hafta","Devre","Ev Sahibi","Ev Gol","Dep Gol","Deplasman","Durum"]
    fk_widths = [6,7,7,22,8,8,22,12]
    for ci, (h, w) in enumerate(zip(fk_hdrs, fk_widths), 1):
        c = ws2.cell(2, ci, h)
        c.font = hfont(9); c.fill = fill(C_SUBHDR)
        c.alignment = ctr(); c.border = border()
        ws2.column_dimensions[get_column_letter(ci)].width = w
    ws2.row_dimensions[2].height = 20

    tum_h = sorted(set(m["week"] for m in fixtures))
    yari  = len(tum_h) // 2

    for ri, m in enumerate(fixtures):
        r = ri + 3
        bg = C_ODD if ri % 2 == 0 else "FFFFFF"
        devre = m.get("devre") or (1 if m["week"] <= yari else 2)
        if m["played"]:
            hg, ag = m["hg"], m["ag"]
            durum = "EV SAHİBİ" if hg > ag else ("DEPLASMAN" if hg < ag else "BERABERE")
            d_color = (C_GREEN_F if hg > ag else (C_RED_F if hg < ag else "B8860B"))
        else:
            hg = ag = ""; durum = "—"; d_color = "888888"

        vals = [m["id"], m["week"], devre, m["home"], hg, ag, m["away"], durum]
        for ci, v in enumerate(vals, 1):
            c = ws2.cell(r, ci, v)
            c.alignment = lft() if ci in (4, 7) else ctr()
            c.border = border(); c.fill = fill(bg)
            c.font = cfont(9, color=(d_color if ci == 8 else "000000"),
                           bold=(ci == 8 and m["played"]))
        ws2.row_dimensions[r].height = 14

    ws2.sheet_properties.tabColor = "375623"

    # ══════════════════════════════════════════════════════════════════════
    # SAYFA 3 — İSTATİSTİKLER
    # ══════════════════════════════════════════════════════════════════════
    ws3 = wb.create_sheet("İstatistikler")
    ws3.sheet_view.showGridLines = False
    ws3.freeze_panes = "A3"

    ws3.merge_cells("A1:O1")
    c = ws3["A1"]
    c.value = "📊  TAKIM İSTATİSTİKLERİ"
    c.font = Font(name="Arial", bold=True, size=14, color="FFFFFF")
    c.fill = fill(C_HEADER); c.alignment = ctr()
    ws3.row_dimensions[1].height = 28

    stat_hdrs = [
        "Sıra","Takım",
        "İç Oyn","İç Gal","İç Ber","İç Mağ","İç A Gol","İç Y Gol","İç Aver",
        "Dış Oyn","Dış Gal","Dış Ber","Dış Mağ","Dış A Gol","Dış Y Gol","Dış Aver",
    ]
    stat_widths = [6,22,8,7,7,7,8,8,8,8,7,7,7,8,8,8]
    grp_s2 = 3; grp_s3 = grp_s2 + 7  # İç=3-9, Dış=10-16

    for ci, (h, w) in enumerate(zip(stat_hdrs, stat_widths), 1):
        c = ws3.cell(2, ci, h)
        is_grp = ci in (grp_s2, grp_s3)
        c.font  = Font(name="Arial", bold=True, color="FFFFFF", size=9)
        c.fill  = fill(C_GRP if is_grp else C_SUBHDR)
        c.alignment = ctr(); c.border = border()
        ws3.column_dimensions[get_column_letter(ci)].width = w
    ws3.row_dimensions[2].height = 20

    for ri, row in enumerate(table):
        r = ri + 3
        bg = C_ODD if ri % 2 == 0 else "FFFFFF"
        s  = row["Sıra"]
        ic_av  = row["İç AG"]  - row["İç YG"]
        dis_av = row["Dış AG"] - row["Dış YG"]
        vals = [
            s, row["Takım"],
            row["İç O"], row["İç G"], row["İç B"], row["İç M"], row["İç AG"], row["İç YG"], ic_av,
            row["Dış O"], row["Dış G"], row["Dış B"], row["Dış M"], row["Dış AG"], row["Dış YG"], dis_av,
        ]
        for ci, v in enumerate(vals, 1):
            c = ws3.cell(r, ci, v)
            c.alignment = lft() if ci == 2 else ctr()
            c.border = grp_border() if ci in (grp_s2, grp_s3) else border()
            if s <= 3:
                c.fill = fill(C_GREEN); c.font = Font(name="Arial", size=10, color=C_GREEN_F)
            elif s >= len(table) - 2:
                c.fill = fill(C_RED);   c.font = Font(name="Arial", size=10, color=C_RED_F)
            elif 4 <= s <= 6:
                c.fill = fill(C_YELLOW); c.font = cfont()
            else:
                c.fill = fill(bg);       c.font = cfont()
        ws3.row_dimensions[r].height = 16

    ws3.sheet_properties.tabColor = "843C0C"

    # Kaydet
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def form_html(form5):
    colors = {"G": "#27ae60", "B": "#f39c12", "M": "#e74c3c"}
    badges = ""
    for item in form5:
        if isinstance(item, tuple) and len(item) == 3:
            f, skor, rakip = item
            tooltip = f'title="{rakip} | {skor}"'
        elif isinstance(item, tuple) and len(item) == 2:
            f, skor = item
            tooltip = f'title="{skor}"'
        else:
            f, tooltip = item, ""
        badges += (
            f'<span {tooltip} style="background:{colors[f]};color:#fff;border-radius:4px;'
            f'padding:1px 6px;margin:1px;font-size:13px;font-weight:700;cursor:default">{f}</span>'
        )
    return badges or '<span style="color:#4a5568">—</span>'

# ── CSS ───────────────────────────────────────────────────────────────────────
st.markdown("""
<style>
  @import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;600;700&display=swap');
  html, body, [class*="css"] { font-family: 'Inter', sans-serif; }
  .lig-table { width:100%; border-collapse:collapse; font-size:14px; }
  .lig-table th {
    background:#1a2340; color:#a0aec0; padding:8px 10px;
    text-align:center; border-bottom:2px solid #2d3748;
    font-weight:600; font-size:12px;
  }
  .lig-table th.left { text-align:left; }
  .lig-table td { padding:7px 10px; text-align:center; border-bottom:1px solid #1e2535; }
  .lig-table td.left { text-align:left; font-weight:600; }
  .lig-table tr:hover td { background:#666666 !important; }
  .zona-cl   { border-left:3px solid #3498db; }
  .zona-al   { border-left:3px solid #9b59b6; }
  .zona-kl   { border-left:3px solid #1abc9c; }
  .zona-kd   { border-left:3px solid #e74c3c; }
  .zona-norm { border-left:3px solid transparent; }
  .sira-badge {
    display:inline-block; width:22px; height:22px; line-height:22px;
    border-radius:50%; font-size:11px; font-weight:700; text-align:center;
  }
  .sira-cl  { background:#3498db; color:#fff; }
  .sira-al  { background:#9b59b6; color:#fff; }
  .sira-kl  { background:#1abc9c; color:#fff; }
  .sira-kd  { background:#e74c3c; color:#fff; }
  .sira-norm{ background:#2d3748; color:#a0aec0; }
  .puan { font-weight:700; font-size:14px; color:#fff; }
  .mac-kart {
    background:#1a2340; border-radius:10px; padding:14px 18px;
    margin-bottom:8px; border:1px solid #2d3748;
  }
  .mac-baslik { color:#718096; font-size:11px; font-weight:600; margin-bottom:6px; }
  .takim-adi { font-size:18px; font-weight:700; color:#fff; text-align:center; }
  .vs { color:#4a5568; font-size:18px; font-weight:700; text-align:center; }
  .metric-card {
    background:#1a2340; border-radius:10px; padding:16px;
    border:1px solid #2d3748; text-align:center; margin-bottom:8px;
  }
  .metric-val { font-size:28px; font-weight:700; color:#fff; }
  .metric-lbl { font-size:12px; color:#718096; margin-top:4px; }
  .sidebar-title { color:#a0aec0; font-size:11px; font-weight:700;
    text-transform:uppercase; letter-spacing:1px; margin-bottom:8px; }
  .ekle-kart {
    background:#1a2340; border-radius:10px; padding:16px;
    border:1px solid #2d3748; margin-bottom:12px;
  }
  .lig-table td.grp-start, .lig-table th.grp-start {
    border-left: 3px solid #718096 !important;
  }
  .lig-table th.puan-hdr { color: #ffd700; }
  div[data-testid="stNumberInput"] input { text-align:center; }
</style>
""", unsafe_allow_html=True)

# ── Başlat ───────────────────────────────────────────────────────────────────
init_db()
migrate_from_json()

teams_list    = db_get_teams()
fixtures_list = db_get_fixtures()
data = make_data(teams_list, fixtures_list)
N = len(data["teams"])

# ════════════════════════════════════════════════════════════════════════════
# SIDEBAR
# ════════════════════════════════════════════════════════════════════════════
with st.sidebar:
    st.markdown("## ⚙️ Lig Ayarları")
    st.markdown("---")

    # Zona ayarları
    st.markdown('<div class="sidebar-title">🏅 Avrupa Zonaları</div>', unsafe_allow_html=True)
    cl_limit = st.number_input("🔵 Şampiyonlar Ligi", min_value=0, max_value=max(N,1), value=min(3, N), step=1, key="cl")
    al_limit = st.number_input("🟣 Avrupa Ligi",       min_value=0, max_value=max(N,1), value=min(5, N), step=1, key="al")
    kl_limit = st.number_input("🟢 Konferans Ligi",    min_value=0, max_value=max(N,1), value=min(7, N), step=1, key="kl")

    st.markdown('<div class="sidebar-title" style="margin-top:10px">🔻 Küme Düşme</div>', unsafe_allow_html=True)
    kd_count = st.number_input("Son kaç takım düşer?", min_value=0, max_value=max(N//2,1), value=min(3, max(N//2,0)), step=1, key="kd")
    kd_start  = N - kd_count + 1

    st.markdown("---")

    # Lig durumu
    played_sb = sum(1 for m in data["fixtures"] if m["played"])
    total_sb  = len(data["fixtures"])
    pct = int(played_sb / total_sb * 100) if total_sb else 0
    st.markdown('<div class="sidebar-title">📊 Lig Durumu</div>', unsafe_allow_html=True)
    st.progress(pct / 100)
    st.markdown(
        f"<div style='color:#718096;font-size:12px;text-align:center'>"
        f"{played_sb} / {total_sb} maç oynandı ({pct}%)</div>",
        unsafe_allow_html=True
    )
    st.markdown(
        f"<div style='color:#718096;font-size:12px;text-align:center;margin-top:4px'>"
        f"{N} takım</div>", unsafe_allow_html=True
    )

# ── Başlık ────────────────────────────────────────────────────────────────────
st.markdown(f"""
<div style="text-align:center;padding:16px 0 8px">
  <span style="font-size:32px">⚽</span>
  <h1 style="margin:4px 0;font-size:26px;color:#fff">LİG YÖNETİCİSİ</h1>
  <p style="color:#718096;font-size:13px">{N} Takım · {total_sb} Maç</p>
</div>
""", unsafe_allow_html=True)

tab1, tab2, tab3, tab4, tab5, tab6 = st.tabs([
    "🏆 Puan Tablosu",
    "⚽ Skor Girişi",
    "📅 Fikstür",
    "➕ Fikstür Oluştur",
    "📈 Analiz",
    "⚙️ Ayarlar"
])

# ════════════════════════════════════════════════════════════════════════════
# TAB 1 — PUAN TABLOSU
# ════════════════════════════════════════════════════════════════════════════
with tab1:
    table = compute_table(data)

    if not table:
        st.info("Henüz takım yok. **Ayarlar** sekmesinden takım ekle.")
    else:
        played = sum(1 for m in data["fixtures"] if m["played"])
        total  = len(data["fixtures"])
        goals  = sum((m["hg"] or 0) + (m["ag"] or 0) for m in data["fixtures"] if m["played"])
        avg_g  = round(goals / played, 2) if played else 0

        c1, c2, c3, c4 = st.columns(4)
        for col, val, lbl in [
            (c1, played, "Oynanan Maç"),
            (c2, total - played, "Kalan Maç"),
            (c3, goals, "Toplam Gol"),
            (c4, avg_g, "Maç Başı Gol"),
        ]:
            col.markdown(f'<div class="metric-card"><div class="metric-val">{val}</div>'
                         f'<div class="metric-lbl">{lbl}</div></div>', unsafe_allow_html=True)

        st.markdown("<br>", unsafe_allow_html=True)

        # Zona açıklaması
        zona_info = []
        if cl_limit > 0:
            zona_info.append(f'<span><span style="color:#3498db">●</span> 1-{cl_limit} Şampiyonlar Ligi</span>')
        if al_limit > cl_limit:
            zona_info.append(f'<span><span style="color:#9b59b6">●</span> {cl_limit+1}-{al_limit} Avrupa Ligi</span>')
        if kl_limit > al_limit:
            zona_info.append(f'<span><span style="color:#1abc9c">●</span> {al_limit+1}-{kl_limit} Konferans Ligi</span>')
        if kd_count > 0:
            zona_info.append(f'<span><span style="color:#e74c3c">●</span> {kd_start}-{N} Küme Düşme</span>')

        if zona_info:
            st.markdown(
                f'<div style="display:flex;gap:16px;margin-bottom:12px;font-size:12px;'
                f'color:#a0aec0;flex-wrap:wrap">{"".join(zona_info)}</div>',
                unsafe_allow_html=True
            )

        html = '<table class="lig-table"><thead><tr>'
        headers = [
            ("Sıra", ""), ("Takım", "left"),
            ("Oyn",""), ("Gal",""), ("Ber",""), ("Mağ",""), ("A Gol",""), ("Y Gol",""), ("Aver",""), ("Puan","puan-hdr"),
            ("İç Oyn","grp"), ("İç Gal",""), ("İç Ber",""), ("İç Mağ",""), ("İç A Gol",""), ("İç Y Gol",""),
            ("Dış Oyn","grp"), ("Dış Gal",""), ("Dış Ber",""), ("Dış Mağ",""), ("Dış A Gol",""), ("Dış Y Gol",""),
            ("Son 5",""),
        ]
        for h, cls in headers:
            if cls == "grp":
                html += f'<th style="border-left:3px solid #718096;text-align:center">{h}</th>'
            elif cls == "left":
                html += f'<th class="left">{h}</th>'
            elif cls == "puan-hdr":
                html += f'<th style="color:#ffd700;text-align:center">{h}</th>'
            else:
                html += f'<th>{h}</th>'
        html += "</tr></thead><tbody>"

        for row in table:
            s = row["Sıra"]
            if cl_limit > 0 and s <= cl_limit:
                zona_cls, sira_cls = "zona-cl", "sira-cl"
            elif al_limit > cl_limit and s <= al_limit:
                zona_cls, sira_cls = "zona-al", "sira-al"
            elif kl_limit > al_limit and s <= kl_limit:
                zona_cls, sira_cls = "zona-kl", "sira-kl"
            elif kd_count > 0 and s >= kd_start:
                zona_cls, sira_cls = "zona-kd", "sira-kd"
            else:
                zona_cls, sira_cls = "zona-norm", "sira-norm"

            av = row["AV"]
            av_color = "#27ae60" if av > 0 else ("#e74c3c" if av < 0 else "#a0aec0")

            html += (
                f'<tr>'
                f'<td class="{zona_cls}"><span class="sira-badge {sira_cls}">{s}</span></td>'
                f'<td class="left">{row["Takım"]}</td>'
                f'<td>{row["O"]}</td><td>{row["G"]}</td><td>{row["B"]}</td><td>{row["M"]}</td>'
                f'<td>{row["AG"]}</td><td>{row["YG"]}</td>'
                f'<td style="color:{av_color};font-weight:600">{av:+d}</td>'
                f'<td class="puan">{row["Puan"]}</td>'
                f'<td style="border-left:3px solid #718096">{row["İç O"]}</td><td>{row["İç G"]}</td><td>{row["İç B"]}</td><td>{row["İç M"]}</td>'
                f'<td>{row["İç AG"]}</td><td>{row["İç YG"]}</td>'
                f'<td style="border-left:3px solid #718096">{row["Dış O"]}</td><td>{row["Dış G"]}</td><td>{row["Dış B"]}</td><td>{row["Dış M"]}</td>'
                f'<td>{row["Dış AG"]}</td><td>{row["Dış YG"]}</td>'
                f'<td>{form_html(row["form"])}</td>'
                f'</tr>'
            )
        html += "</tbody></table>"
        st.markdown(html, unsafe_allow_html=True)

# ════════════════════════════════════════════════════════════════════════════
# TAB 2 — SKOR GİRİŞİ
# ════════════════════════════════════════════════════════════════════════════
with tab2:
    if not data["fixtures"]:
        st.info("Henüz maç yok. **Fikstür Oluştur** sekmesinden maç ekle.")
    else:
        all_weeks = sorted(set(m["week"] for m in data["fixtures"]))

        col_l, col_r = st.columns([3, 1])
        col_l.markdown("### Skor Girişi")
        week = col_r.selectbox("Hafta", all_weeks,
                                format_func=lambda w: f"Hafta {w}", key="week_sel")

        week_matches = [m for m in data["fixtures"] if m["week"] == week]
        f1, f2 = st.columns(2)
        show_played   = f1.checkbox("Oynanmışları göster", value=True)
        show_unplayed = f2.checkbox("Oynanmamışları göster", value=True)

        filtered = [m for m in week_matches
                    if (m["played"] and show_played) or (not m["played"] and show_unplayed)]

        played_w = sum(1 for m in week_matches if m["played"])
        st.markdown(
            f"<p style='color:#718096;font-size:13px'>"
            f"{week}. Hafta · {len(week_matches)} maç · {played_w} oynandı</p>",
            unsafe_allow_html=True
        )

        changed = False
        for m in filtered:
            st.markdown(
                f'<div class="mac-kart">'
                f'<div class="mac-baslik">MAÇ {m["id"]} · Hafta {m["week"]}</div>',
                unsafe_allow_html=True
            )
            c1, c2, c3, c4, c5 = st.columns([3, 1, 1, 1, 3])
            c1.markdown(f'<div class="takim-adi">{m["home"]}</div>', unsafe_allow_html=True)
            c3.markdown('<div class="vs">—</div>', unsafe_allow_html=True)
            c5.markdown(f'<div class="takim-adi">{m["away"]}</div>', unsafe_allow_html=True)

            hg_new = c2.number_input("", min_value=0, max_value=30,
                                      value=m["hg"] or 0, key=f"hg_{m['id']}",
                                      label_visibility="collapsed")
            ag_new = c4.number_input("", min_value=0, max_value=30,
                                      value=m["ag"] or 0, key=f"ag_{m['id']}",
                                      label_visibility="collapsed")
            played_new = st.checkbox("Oynandı ✓", value=m["played"], key=f"p_{m['id']}")
            st.markdown("</div>", unsafe_allow_html=True)

            if hg_new != m["hg"] or ag_new != m["ag"] or played_new != m["played"]:
                m["hg"] = hg_new; m["ag"] = ag_new; m["played"] = played_new
                changed = True

        if changed:
            # SQLite'a yaz
            for m in filtered:
                db_update_score(m["id"], m["hg"], m["ag"], m["played"])
            st.rerun()

        if st.button("💾 Kaydet", use_container_width=True, type="primary", key="skor_kaydet"):
            for m in filtered:
                db_update_score(m["id"], m["hg"], m["ag"], m["played"])
            st.success("Kaydedildi!")

# ════════════════════════════════════════════════════════════════════════════
# TAB 3 — FİKSTÜR (görüntüle + sil)
# ════════════════════════════════════════════════════════════════════════════
with tab3:
    st.markdown("### 📅 Fikstür")

    if not data["fixtures"]:
        st.info("Henüz maç yok. **Fikstür Oluştur** sekmesinden maç ekle.")
    else:
        fa, fb, fc = st.columns(3)
        team_f = fa.selectbox("Takım", ["Tümü"] + data["teams"], key="fiks_team")
        hafta_options = ["Tümü"] + sorted(set(m["week"] for m in data["fixtures"]))
        hafta_f = fb.selectbox("Hafta", hafta_options, key="fiks_hafta")
        stat_f  = fc.selectbox("Durum", ["Tümü", "Oynandı", "Oynanmadı"], key="fiks_stat")

        flt = list(data["fixtures"])
        if team_f != "Tümü":
            flt = [m for m in flt if m["home"] == team_f or m["away"] == team_f]
        if hafta_f != "Tümü":
            flt = [m for m in flt if m["week"] == hafta_f]
        if stat_f == "Oynandı":
            flt = [m for m in flt if m["played"]]
        elif stat_f == "Oynanmadı":
            flt = [m for m in flt if not m["played"]]

        # Takım istatistik paneli
        if team_f != "Tümü":
            table_all = compute_table(data)
            t_row = next((r for r in table_all if r["Takım"] == team_f), None)
            if t_row:
                st.markdown(f"#### 📊 {team_f} İstatistikleri")
                sc1,sc2,sc3,sc4,sc5,sc6,sc7,sc8 = st.columns(8)
                for col, val, lbl in [
                    (sc1, t_row["Sıra"],  "Sıra"),
                    (sc2, t_row["O"],     "Oyn"),
                    (sc3, t_row["G"],     "Gal"),
                    (sc4, t_row["B"],     "Ber"),
                    (sc5, t_row["M"],     "Mağ"),
                    (sc6, f'{t_row["AV"]:+d}', "Aver"),
                    (sc7, t_row["Puan"], "Puan"),
                    (sc8, f'{t_row["AG"]}-{t_row["YG"]}', "A-Y Gol"),
                ]:
                    col.markdown(
                        f'<div class="metric-card"><div class="metric-val">{val}</div>'
                        f'<div class="metric-lbl">{lbl}</div></div>',
                        unsafe_allow_html=True
                    )
                # İç/Dış özet
                ic1,ic2,ic3,ic4 = st.columns(4)
                for col, val, lbl in [
                    (ic1, f'{t_row["İç G"]}/{t_row["İç B"]}/{t_row["İç M"]}', "İç G/B/M"),
                    (ic2, f'{t_row["İç AG"]}-{t_row["İç YG"]}',               "İç A-Y Gol"),
                    (ic3, f'{t_row["Dış G"]}/{t_row["Dış B"]}/{t_row["Dış M"]}', "Dış G/B/M"),
                    (ic4, f'{t_row["Dış AG"]}-{t_row["Dış YG"]}',             "Dış A-Y Gol"),
                ]:
                    col.markdown(
                        f'<div class="metric-card"><div class="metric-val" style="font-size:18px">{val}</div>'
                        f'<div class="metric-lbl">{lbl}</div></div>',
                        unsafe_allow_html=True
                    )
                st.markdown(f"**Son 5:** {form_html(t_row['form'])}", unsafe_allow_html=True)
                st.markdown("---")

        st.markdown(f"<p style='color:#718096;font-size:13px'>{len(flt)} maç</p>",
                    unsafe_allow_html=True)

        tum_haftalar = sorted(set(m["week"] for m in data["fixtures"]))
        yari_hafta   = len(tum_haftalar) // 2

        html2 = '<table class="lig-table"><thead><tr>'
        for h in ["No", "Hafta", "Devre", "Ev Sahibi", "Skor", "Deplasman", "Durum", "Sil"]:
            cls = "left" if h in ["Ev Sahibi", "Deplasman"] else ""
            html2 += f'<th class="{cls}">{h}</th>'
        html2 += "</tr></thead><tbody>"

        for m in flt:
            if m["played"]:
                hg, ag = m["hg"], m["ag"]
                skor = f'<b>{hg} - {ag}</b>'
                if hg > ag:
                    res_color, res = "#27ae60", "EV SAHİBİ"
                    home_style = "font-weight:700"
                    away_style = "opacity:0.45"
                elif hg < ag:
                    res_color, res = "#e74c3c", "DEPLASMAN"
                    home_style = "opacity:0.45"
                    away_style = "font-weight:700"
                else:
                    res_color, res = "#f39c12", "BERABERE"
                    home_style = away_style = "font-weight:600"
                durum = f'<span style="color:{res_color};font-size:11px;font-weight:700">{res}</span>'
            else:
                skor  = '<span style="color:#4a5568">vs</span>'
                durum = '<span style="color:#4a5568;font-size:11px">—</span>'
                home_style = away_style = ""

            devre = m.get("devre") or (1 if m["week"] <= yari_hafta else 2)
            devre_color = "#3498db" if devre == 1 else "#9b59b6"
            devre_html  = f'<span style="color:{devre_color};font-weight:700">{devre}</span>'

            html2 += (
                f'<tr><td>{m["id"]}</td><td>{m["week"]}</td><td>{devre_html}</td>'
                f'<td class="left" style="{home_style}">{m["home"]}</td>'
                f'<td>{skor}</td>'
                f'<td class="left" style="{away_style}">{m["away"]}</td>'
                f'<td>{durum}</td>'
                f'<td>—</td></tr>'
            )
        html2 += "</tbody></table>"
        st.markdown(html2, unsafe_allow_html=True)

        st.markdown("<br>", unsafe_allow_html=True)
        with st.expander("🗑️ Maç Sil"):
            mac_ids = [m["id"] for m in data["fixtures"]]
            mac_labels = {
                m["id"]: f"#{m['id']} | H{m['week']} | {m['home']} vs {m['away']}"
                for m in data["fixtures"]
            }
            del_id = st.selectbox("Silinecek maç", mac_ids,
                                   format_func=lambda x: mac_labels[x], key="del_mac")
            if st.button("🗑️ Seçili Maçı Sil", type="secondary", key="del_mac_btn"):
                db_delete_fixture(del_id)
                st.success("Maç silindi.")
                st.rerun()

# ════════════════════════════════════════════════════════════════════════════
# TAB 4 — FİKSTÜR OLUŞTUR
# ════════════════════════════════════════════════════════════════════════════
with tab4:
    st.markdown("### ➕ Fikstür Oluştur")

    if not data["teams"]:
        st.warning("Önce **Ayarlar** sekmesinden takımları ekle.")
    else:
        mod = st.radio("Mod seç", ["Otomatik Round-Robin", "Tek Maç Ekle", "Hafta Oluştur"], horizontal=True)

        st.markdown("---")

        teams = data["teams"]

        # ── MOD 0: Otomatik Round-Robin ───────────────────────────────────
        if mod == "Otomatik Round-Robin":

            n = len(teams)
            if n % 2 != 0:
                st.warning("Round-robin için çift sayıda takım gerekiyor. Ayarlar'dan bir takım daha ekle.", icon="⚠️")
            else:
                total_mac = n * (n - 1)
                total_hafta = (n - 1) * 2
                st.info(f"**{n} takım** → {total_mac} maç · {total_hafta} hafta (çift devreli)", icon="📋")

                mevcut = len(data["fixtures"])
                if mevcut > 0:
                    st.warning(f"Mevcut {mevcut} maç silinecek ve yeni fikstür oluşturulacak!", icon="⚠️")

                if st.button("🔄 Fikstürü Oluştur", type="primary", use_container_width=True, key="rr_btn"):
                    def round_robin(teams):
                        n = len(teams)
                        lst = list(range(n))
                        rounds = []
                        for _ in range(n - 1):
                            pairs = [(lst[i], lst[n - 1 - i]) for i in range(n // 2)]
                            rounds.append(pairs)
                            lst = [lst[0]] + [lst[-1]] + lst[1:-1]
                        return rounds

                    r1 = round_robin(teams)
                    r2 = [[(b, a) for a, b in rnd] for rnd in r1]
                    rows = []
                    for devre, rounds in enumerate([r1, r2], 1):
                        for w_idx, pairs in enumerate(rounds):
                            week = w_idx + 1 + (n - 1) * (devre - 1)
                            for hi, ai in pairs:
                                rows.append((week, devre, teams[hi], teams[ai]))
                    db_clear_fixtures()
                    db_bulk_add_fixtures(rows)
                    st.success(f"✅ {len(rows)} maçlık fikstür oluşturuldu!")
                    st.rerun()

        # ── MOD 1: Tek Maç Ekle ───────────────────────────────────────────
        elif mod == "Tek Maç Ekle":
            st.markdown("#### Maç Bilgileri")
            with st.container():
                c1, c2, c3 = st.columns([2, 1, 2])
                home = c1.selectbox("Ev Sahibi", teams, key="add_home")
                away_teams = [t for t in teams if t != home]
                away = c3.selectbox("Deplasman", away_teams, key="add_away")

                hafta_mevcut = sorted(set(m["week"] for m in data["fixtures"]))
                max_hafta = max(hafta_mevcut, default=0)

                c4, c5, c6 = st.columns(3)
                hafta_mod = c4.radio("Hafta", ["Var olan", "Yeni"], horizontal=True, key="hafta_mod")
                if hafta_mod == "Var olan" and hafta_mevcut:
                    hafta = c5.selectbox("Hafta seç", hafta_mevcut, key="add_hafta_sel")
                else:
                    hafta = c5.number_input("Hafta no", min_value=1,
                                             value=max_hafta + 1, step=1, key="add_hafta_new")

                # Skor şimdi girilsin mi?
                skor_simdi = c6.checkbox("Skoru şimdi gir", key="skor_simdi")

                hg_val, ag_val, played_val = 0, 0, False
                if skor_simdi:
                    cs1, cs2, cs3 = st.columns([1,1,1])
                    hg_val = cs1.number_input(f"{home} gol", min_value=0, max_value=30,
                                               value=0, key="add_hg")
                    ag_val = cs3.number_input(f"{away} gol", min_value=0, max_value=30,
                                               value=0, key="add_ag")
                    played_val = True

                if st.button("➕ Maç Ekle", type="primary", use_container_width=True, key="ekle_mac"):
                    if home == away:
                        st.error("Ev sahibi ve deplasman aynı olamaz.")
                    else:
                        devre = get_devre(int(hafta), data["fixtures"])
                        db_add_fixture(int(hafta), devre, home, away,
                                       hg_val if played_val else None,
                                       ag_val if played_val else None,
                                       played_val)
                        st.success(f"✅ {home} vs {away} — Hafta {hafta} eklendi!")
                        st.rerun()

        # ── MOD 2: Hafta Oluştur ──────────────────────────────────────────
        elif mod == "Hafta Oluştur":
            st.markdown("#### Hafta Oluştur")

            hafta_mevcut = sorted(set(m["week"] for m in data["fixtures"]))
            max_hafta = max(hafta_mevcut, default=0)

            hc1, hc2 = st.columns(2)
            hafta_mod2 = hc1.radio("Hafta", ["Mevcut hafta", "Yeni hafta"],
                                    horizontal=True, key="hafta_mod2")

            if hafta_mod2 == "Mevcut hafta":
                if not hafta_mevcut:
                    st.info("Henüz hiç hafta yok. 'Yeni hafta' seç.")
                    hafta_no = 1
                    hafta_maclar = []
                else:
                    hafta_no = hc2.selectbox(
                        "Hafta seç", hafta_mevcut,
                        format_func=lambda w: f"Hafta {w}",
                        key="hw_sel"
                    )
                    hafta_maclar = [m for m in data["fixtures"] if m["week"] == hafta_no]
            else:
                hafta_no = hc2.number_input("Hafta no", min_value=1,
                                             value=max_hafta + 1, step=1, key="hw_new")
                hafta_maclar = []

            st.markdown("---")

            # Mevcut maçlar varsa → selectbox'lar onlarla dolu gelsin
            # Yoksa boş form göster
            if hafta_maclar:
                st.markdown(f"**Hafta {hafta_no} — {len(hafta_maclar)} maç** (düzenle veya yeni ekle)")

                # Önce tablo göster
                tbl = '<table class="lig-table"><thead><tr>'
                for h in ["No", "Ev Sahibi", "Skor", "Deplasman", "Durum"]:
                    cls = "left" if h in ["Ev Sahibi", "Deplasman"] else ""
                    tbl += f'<th class="{cls}">{h}</th>'
                tbl += "</tr></thead><tbody>"
                for m in hafta_maclar:
                    if m["played"]:
                        skor = f"<b>{m['hg']} - {m['ag']}</b>"
                        if m["hg"] > m["ag"]:   dc, dr = "#27ae60", "EV"
                        elif m["hg"] < m["ag"]: dc, dr = "#e74c3c", "DEP"
                        else:                   dc, dr = "#f39c12", "BER"
                        durum = f'<span style="color:{dc};font-weight:700;font-size:11px">{dr}</span>'
                    else:
                        skor  = '<span style="color:#4a5568">vs</span>'
                        durum = '<span style="color:#4a5568;font-size:11px">—</span>'
                    tbl += (f"<tr><td>{m['id']}</td>"
                            f'<td class="left">{m["home"]}</td>'
                            f"<td>{skor}</td>"
                            f'<td class="left">{m["away"]}</td>'
                            f"<td>{durum}</td></tr>")
                tbl += "</tbody></table>"
                st.markdown(tbl, unsafe_allow_html=True)
                st.markdown("<br>", unsafe_allow_html=True)

                n_mac = len(hafta_maclar)
            else:
                n_mac = st.number_input("Kaç maç ekleyeceksin?", min_value=1,
                                         max_value=max(len(teams)//2, 1),
                                         value=min(9, max(len(teams)//2, 1)),
                                         step=1, key="n_mac_hafta")
                n_mac = int(n_mac)

            mac_listesi = []
            valid = True

            for i in range(n_mac):
                # Mevcut maçtan varsayılan değerleri al
                if i < len(hafta_maclar):
                    default_ev  = hafta_maclar[i]["home"]
                    default_dep = hafta_maclar[i]["away"]
                    mac_id      = hafta_maclar[i]["id"]
                    mac_label   = f"**Maç {i+1}** (#{mac_id})"
                else:
                    default_ev  = teams[0]
                    default_dep = teams[1] if len(teams) > 1 else teams[0]
                    mac_label   = f"**Maç {i+1}** (yeni)"

                # teams listesinde index bul
                ev_idx  = teams.index(default_ev)  if default_ev  in teams else 0
                dep_idx = teams.index(default_dep) if default_dep in teams else 0

                st.markdown(mac_label)
                mc1, mc2, mc3 = st.columns([2, 1, 2])
                ev  = mc1.selectbox("Ev",  teams, index=ev_idx,  key=f"hw_h_{i}")
                dep = mc3.selectbox("Dep", teams, index=dep_idx, key=f"hw_a_{i}")

                if ev == dep:
                    mc2.markdown("<br>", unsafe_allow_html=True)
                    st.warning(f"Maç {i+1}: Ev ve deplasman aynı olamaz.", icon="⚠️")
                    valid = False

                mac_listesi.append((i, ev, dep))

            st.markdown("<br>", unsafe_allow_html=True)

            btn_label = "💾 Haftayı Güncelle" if hafta_maclar else "➕ Haftayı Ekle"
            if st.button(btn_label, type="primary",
                          use_container_width=True, key="ekle_hafta"):
                if not valid:
                    st.error("Lütfen hataları düzelt.")
                else:
                    if hafta_maclar:
                        # Mevcut maçları güncelle
                        for i, ev, dep in mac_listesi:
                            if i < len(hafta_maclar):
                                db_update_fixture_teams(hafta_maclar[i]["id"], ev, dep)
                        st.success(f"✅ Hafta {hafta_no} güncellendi!")
                    else:
                        devre = get_devre(int(hafta_no), data["fixtures"])
                        rows = [(int(hafta_no), devre, ev, dep) for _, ev, dep in mac_listesi]
                        db_bulk_add_fixtures(rows)
                        st.success(f"✅ {len(rows)} maç — Hafta {hafta_no}'e eklendi!")
                    st.rerun()

# ════════════════════════════════════════════════════════════════════════════
# TAB 5 — ANALİZ
# ════════════════════════════════════════════════════════════════════════════
with tab5:
    st.markdown("### 📈 Analiz")

    if not data["fixtures"] or not any(m["played"] for m in data["fixtures"]):
        st.info("Henüz oynanan maç yok. Skor girişi yaptıktan sonra analiz görünür.")
    else:
        played_fixtures = [m for m in data["fixtures"] if m["played"]]
        table = compute_table(data)

        analiz_tab1, analiz_tab2, analiz_tab3, analiz_tab4 = st.tabs([
            "📈 Puan Grafiği", "⚽ Gol İstatistikleri", "🆚 Kafa Kafaya", "📋 Haftalık Durum"
        ])

        # ── ANALİZ TAB 1: Hafta hafta puan grafiği ───────────────────────
        with analiz_tab1:
            st.markdown("#### Hafta Hafta Puan Seyri")

            all_teams = data["teams"]
            selected_teams = st.multiselect(
                "Takımları seç", all_teams,
                default=all_teams[:5] if len(all_teams) >= 5 else all_teams,
                key="grafik_teams"
            )

            if not selected_teams:
                st.info("En az bir takım seç.")
            else:
                # Her hafta sonunda kümülatif puan hesapla
                all_weeks = sorted(set(m["week"] for m in played_fixtures))

                # Takım başına haftalık kümülatif puan
                team_weekly = {t: [] for t in selected_teams}
                team_cumulative = {t: 0 for t in all_teams}

                weekly_data = []
                for week in all_weeks:
                    week_matches = [m for m in played_fixtures if m["week"] == week]
                    for m in week_matches:
                        h, a, hg, ag = m["home"], m["away"], m["hg"], m["ag"]
                        if hg > ag:
                            if h in team_cumulative: team_cumulative[h] += 3
                        elif hg < ag:
                            if a in team_cumulative: team_cumulative[a] += 3
                        else:
                            if h in team_cumulative: team_cumulative[h] += 1
                            if a in team_cumulative: team_cumulative[a] += 1
                    for t in selected_teams:
                        team_weekly[t].append(team_cumulative[t])

                # Plotly grafik
                try:
                    import plotly.graph_objects as go
                    fig = go.Figure()
                    colors_list = [
                        "#3498db","#e74c3c","#2ecc71","#f39c12","#9b59b6",
                        "#1abc9c","#e67e22","#34495e","#e91e63","#00bcd4"
                    ]
                    for idx, t in enumerate(selected_teams):
                        fig.add_trace(go.Scatter(
                            x=list(all_weeks),
                            y=team_weekly[t],
                            mode="lines+markers",
                            name=t,
                            line=dict(color=colors_list[idx % len(colors_list)], width=2),
                            marker=dict(size=6)
                        ))
                    fig.update_layout(
                        xaxis_title="Hafta",
                        yaxis_title="Puan",
                        hovermode="x unified",
                        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1),
                        height=450,
                        margin=dict(l=40, r=20, t=40, b=40),
                        plot_bgcolor="rgba(0,0,0,0)",
                        paper_bgcolor="rgba(0,0,0,0)",
                        font=dict(color="#a0aec0"),
                        xaxis=dict(gridcolor="#2d3748", tickmode="linear", dtick=1),
                        yaxis=dict(gridcolor="#2d3748"),
                    )
                    st.plotly_chart(fig, use_container_width=True)
                except ImportError:
                    # Plotly yoksa basit tablo göster
                    st.warning("Grafik için `pip install plotly` gerekiyor. Tablo olarak gösteriliyor.")
                    rows_data = {"Hafta": all_weeks}
                    for t in selected_teams:
                        rows_data[t] = team_weekly[t]
                    import pandas as pd
                    st.dataframe(pd.DataFrame(rows_data), use_container_width=True)

        # ── ANALİZ TAB 2: Gol istatistikleri ─────────────────────────────
        with analiz_tab2:
            st.markdown("#### Gol İstatistikleri")

            # Genel özet
            total_goals = sum((m["hg"] or 0) + (m["ag"] or 0) for m in played_fixtures)
            total_played = len(played_fixtures)
            home_goals   = sum(m["hg"] or 0 for m in played_fixtures)
            away_goals   = sum(m["ag"] or 0 for m in played_fixtures)

            g1, g2, g3, g4 = st.columns(4)
            for col, val, lbl in [
                (g1, total_goals,  "Toplam Gol"),
                (g2, f"{total_goals/total_played:.2f}" if total_played else 0, "Maç Başı Gol"),
                (g3, home_goals,   "Ev Sahibi Gol"),
                (g4, away_goals,   "Deplasman Gol"),
            ]:
                col.markdown(
                    f'<div class="metric-card"><div class="metric-val">{val}</div>'
                    f'<div class="metric-lbl">{lbl}</div></div>',
                    unsafe_allow_html=True
                )

            st.markdown("<br>", unsafe_allow_html=True)

            # Takım bazlı gol tablosu
            gol_stats = []
            for row in table:
                t = row["Takım"]
                mac_basi = round(row["AG"] / row["O"], 2) if row["O"] > 0 else 0
                yenilen_basi = round(row["YG"] / row["O"], 2) if row["O"] > 0 else 0
                gol_stats.append({
                    "Sıra": row["Sıra"],
                    "Takım": t,
                    "Att": row["AG"],
                    "Yedi": row["YG"],
                    "Fark": row["AV"],
                    "Maç Başı Att": mac_basi,
                    "Maç Başı Yedi": yenilen_basi,
                    "İç Att": row["İç AG"],
                    "İç Yedi": row["İç YG"],
                    "Dış Att": row["Dış AG"],
                    "Dış Yedi": row["Dış YG"],
                })

            # En çok gol atan sıralaması
            col_a, col_b = st.columns(2)

            with col_a:
                st.markdown("**⬆️ En Çok Gol Atan**")
                sorted_att = sorted(gol_stats, key=lambda x: -x["Att"])
                html_g = '<table class="lig-table"><thead><tr>'
                for h in ["Takım", "Att", "M Başı", "İç", "Dış"]:
                    html_g += f'<th class="{"left" if h=="Takım" else ""}">{h}</th>'
                html_g += "</tr></thead><tbody>"
                for i, r in enumerate(sorted_att):
                    bg = "#bcbcbc" if i % 2 == 0 else ""
                    html_g += (
                        f'<tr style="background:{bg}">'
                        f'<td class="left">{r["Takım"]}</td>'
                        f'<td style="color:#27ae60;font-weight:700">{r["Att"]}</td>'
                        f'<td>{r["Maç Başı Att"]}</td>'
                        f'<td>{r["İç Att"]}</td>'
                        f'<td>{r["Dış Att"]}</td>'
                        f'</tr>'
                    )
                html_g += "</tbody></table>"
                st.markdown(html_g, unsafe_allow_html=True)

            with col_b:
                st.markdown("**⬇️ En Çok Gol Yiyen**")
                sorted_yedi = sorted(gol_stats, key=lambda x: -x["Yedi"])
                html_y = '<table class="lig-table"><thead><tr>'
                for h in ["Takım", "Yedi", "M Başı", "İç", "Dış"]:
                    html_y += f'<th class="{"left" if h=="Takım" else ""}">{h}</th>'
                html_y += "</tr></thead><tbody>"
                for i, r in enumerate(sorted_yedi):
                    bg = "#1a2340" if i % 2 == 0 else ""
                    html_y += (
                        f'<tr style="background:{bg}">'
                        f'<td class="left">{r["Takım"]}</td>'
                        f'<td style="color:#e74c3c;font-weight:700">{r["Yedi"]}</td>'
                        f'<td>{r["Maç Başı Yedi"]}</td>'
                        f'<td>{r["İç Yedi"]}</td>'
                        f'<td>{r["Dış Yedi"]}</td>'
                        f'</tr>'
                    )
                html_y += "</tbody></table>"
                st.markdown(html_y, unsafe_allow_html=True)

            # Hafta bazlı gol grafiği
            st.markdown("<br>**📊 Hafta Bazlı Gol Ortalaması**")
            all_weeks_g = sorted(set(m["week"] for m in played_fixtures))
            hafta_goller = []
            for w in all_weeks_g:
                wm = [m for m in played_fixtures if m["week"] == w]
                toplam = sum((m["hg"] or 0) + (m["ag"] or 0) for m in wm)
                ort = round(toplam / len(wm), 2) if wm else 0
                hafta_goller.append(ort)

            try:
                import plotly.graph_objects as go
                fig2 = go.Figure()
                fig2.add_trace(go.Bar(
                    x=list(all_weeks_g),
                    y=hafta_goller,
                    marker_color="#3498db",
                    name="Maç Başı Gol"
                ))
                fig2.add_hline(
                    y=total_goals/total_played if total_played else 0,
                    line_dash="dash", line_color="#f39c12",
                    annotation_text=f"Ortalama: {total_goals/total_played:.2f}" if total_played else ""
                )
                fig2.update_layout(
                    xaxis_title="Hafta", yaxis_title="Maç Başı Gol",
                    height=300, margin=dict(l=40,r=20,t=20,b=40),
                    plot_bgcolor="rgba(0,0,0,0)", paper_bgcolor="rgba(0,0,0,0)",
                    font=dict(color="#a0aec0"),
                    xaxis=dict(gridcolor="#2d3748", tickmode="linear", dtick=1),
                    yaxis=dict(gridcolor="#2d3748"),
                )
                st.plotly_chart(fig2, use_container_width=True)
            except ImportError:
                pass

        # ── ANALİZ TAB 3: Kafa kafaya ────────────────────────────────────
        with analiz_tab3:
            st.markdown("#### 🆚 Kafa Kafaya Karşılaştırma")

            teams_list = data["teams"]
            if len(teams_list) < 2:
                st.info("En az 2 takım gerekiyor.")
            else:
                kk1, kk2 = st.columns(2)
                takim1 = kk1.selectbox("1. Takım", teams_list, key="kk_t1")
                takim2 = kk2.selectbox("2. Takım",
                    [t for t in teams_list if t != takim1],
                    key="kk_t2"
                )

                # Aralarındaki maçlar
                h2h = [m for m in played_fixtures if
                       (m["home"] == takim1 and m["away"] == takim2) or
                       (m["home"] == takim2 and m["away"] == takim1)]

                if not h2h:
                    st.info(f"{takim1} ile {takim2} henüz karşılaşmadı.")
                else:
                    # Özet istatistik
                    t1_g = t1_b = t1_m = t1_ag = t1_yg = 0
                    for m in h2h:
                        if m["home"] == takim1:
                            hg, ag = m["hg"], m["ag"]
                        else:
                            hg, ag = m["ag"], m["hg"]
                        t1_ag += hg; t1_yg += ag
                        if hg > ag:   t1_g += 1
                        elif hg < ag: t1_m += 1
                        else:         t1_b += 1
                    t2_g = t1_m; t2_b = t1_b; t2_m = t1_g

                    # Karşılaştırma kartları
                    st.markdown(f"**{len(h2h)} karşılaşma**")
                    kc1, kc2, kc3 = st.columns(3)

                    kc1.markdown(
                        f'<div class="metric-card">'
                        f'<div style="font-size:13px;color:#a0aec0;margin-bottom:8px">{takim1}</div>'
                        f'<div style="font-size:28px;font-weight:700;color:#27ae60">{t1_g}</div>'
                        f'<div style="font-size:11px;color:#718096">Galibiyet</div>'
                        f'<div style="font-size:20px;font-weight:600;margin-top:6px">{t1_ag}-{t1_yg}</div>'
                        f'<div style="font-size:11px;color:#718096">Gol</div>'
                        f'</div>', unsafe_allow_html=True
                    )
                    kc2.markdown(
                        f'<div class="metric-card" style="text-align:center">'
                        f'<div style="font-size:28px;font-weight:700;color:#f39c12">{t1_b}</div>'
                        f'<div style="font-size:11px;color:#718096">Beraberlik</div>'
                        f'</div>', unsafe_allow_html=True
                    )
                    kc3.markdown(
                        f'<div class="metric-card">'
                        f'<div style="font-size:13px;color:#a0aec0;margin-bottom:8px">{takim2}</div>'
                        f'<div style="font-size:28px;font-weight:700;color:#27ae60">{t2_g}</div>'
                        f'<div style="font-size:11px;color:#718096">Galibiyet</div>'
                        f'<div style="font-size:20px;font-weight:600;margin-top:6px">{t1_yg}-{t1_ag}</div>'
                        f'<div style="font-size:11px;color:#718096">Gol</div>'
                        f'</div>', unsafe_allow_html=True
                    )

                    st.markdown("<br>**Karşılaşmalar:**")
                    html_h2h = '<table class="lig-table"><thead><tr>'
                    for h in ["Hafta","Ev Sahibi","Skor","Deplasman","Sonuç"]:
                        cls = "left" if h in ["Ev Sahibi","Deplasman"] else ""
                        html_h2h += f'<th class="{cls}">{h}</th>'
                    html_h2h += "</tr></thead><tbody>"

                    for m in sorted(h2h, key=lambda x: x["week"]):
                        hg, ag = m["hg"], m["ag"]
                        if m["home"] == takim1:
                            kazanan = takim1 if hg > ag else (takim2 if hg < ag else "Berabere")
                        else:
                            kazanan = takim2 if hg > ag else (takim1 if hg < ag else "Berabere")

                        if kazanan == takim1:   kc, kl = "#27ae60", takim1
                        elif kazanan == takim2: kc, kl = "#e74c3c", takim2
                        else:                   kc, kl = "#f39c12", "Berabere"

                        html_h2h += (
                            f'<tr>'
                            f'<td>{m["week"]}</td>'
                            f'<td class="left">{m["home"]}</td>'
                            f'<td><b>{hg} - {ag}</b></td>'
                            f'<td class="left">{m["away"]}</td>'
                            f'<td><span style="color:{kc};font-weight:700">{kl}</span></td>'
                            f'</tr>'
                        )
                    html_h2h += "</tbody></table>"
                    st.markdown(html_h2h, unsafe_allow_html=True)

                    # Genel puan tablosundaki durumları
                    st.markdown("<br>**Lig Sıralaması:**")
                    t1_row = next((r for r in table if r["Takım"] == takim1), None)
                    t2_row = next((r for r in table if r["Takım"] == takim2), None)
                    if t1_row and t2_row:
                        cs1, cs2 = st.columns(2)
                        for col, row in [(cs1, t1_row), (cs2, t2_row)]:
                            col.markdown(
                                f'<div class="metric-card">'
                                f'<div style="font-size:13px;font-weight:700;margin-bottom:8px">{row["Takım"]}</div>'
                                f'<div style="font-size:24px;font-weight:700">{row["Sıra"]}. Sıra</div>'
                                f'<div style="margin-top:6px;font-size:13px">'
                                f'<span style="color:#27ae60">{row["G"]}G</span> '
                                f'<span style="color:#f39c12">{row["B"]}B</span> '
                                f'<span style="color:#e74c3c">{row["M"]}M</span> · '
                                f'<b>{row["Puan"]} puan</b>'
                                f'</div>'
                                f'</div>',
                                unsafe_allow_html=True
                            )

        # ── ANALİZ TAB 4: Haftalık puan durumu ───────────────────────────
        with analiz_tab4:
            st.markdown("#### 📋 Haftalık Puan Durumu")

            all_weeks_s = sorted(set(m["week"] for m in played_fixtures))
            if not all_weeks_s:
                st.info("Henüz oynanan maç yok.")
            else:
                selected_week = st.select_slider(
                    "Hafta seç",
                    options=all_weeks_s,
                    value=all_weeks_s[-1],
                    format_func=lambda w: f"Hafta {w}",
                    key="haftalik_week"
                )

                # O haftaya kadar oynanan maçlardan puan tablosu hesapla
                matches_upto = [m for m in played_fixtures if m["week"] <= selected_week]
                week_matches = [m for m in played_fixtures if m["week"] == selected_week]

                # Geçici data objesi oluştur
                temp_data = {"n": N, "teams": data["teams"],
                             "fixtures": matches_upto + [m for m in data["fixtures"] if not m["played"]]}
                week_table = compute_table({"n": N, "teams": data["teams"], "fixtures": matches_upto})

                # Özet bilgi
                oynanan_w    = len(matches_upto)
                gol_toplam   = sum((m["hg"] or 0) + (m["ag"] or 0) for m in matches_upto)
                gol_bu_hafta = sum((m["hg"] or 0) + (m["ag"] or 0) for m in week_matches)

                w1, w2, w3, w4 = st.columns(4)
                w1.markdown(f'<div class="metric-card"><div class="metric-val">{selected_week}</div><div class="metric-lbl">Hafta</div></div>', unsafe_allow_html=True)
                w2.markdown(f'<div class="metric-card"><div class="metric-val">{len(week_matches)}</div><div class="metric-lbl">{selected_week}. Hafta Maç</div></div>', unsafe_allow_html=True)
                w3.markdown(f'<div class="metric-card"><div class="metric-val">{gol_bu_hafta}</div><div class="metric-lbl">{selected_week}. Hafta Gol</div></div>', unsafe_allow_html=True)
                w4.markdown(f'<div class="metric-card"><div class="metric-val">{gol_toplam}</div><div class="metric-lbl">Toplam Gol (1-{selected_week})</div></div>', unsafe_allow_html=True)

                st.markdown("<br>", unsafe_allow_html=True)

                left_col, right_col = st.columns([3, 2])

                # Sol: Puan tablosu o haftaya kadar
                with left_col:
                    st.markdown(f"**{selected_week}. Hafta Sonrası Puan Durumu**")
                    html_wt = '<table class="lig-table"><thead><tr>'
                    for h in ["Sıra","Takım","O","G","B","M","AG","YG","Aver","Puan"]:
                        cls = "left" if h == "Takım" else ""
                        html_wt += f'<th class="{cls}">{h}</th>'
                    html_wt += "</tr></thead><tbody>"

                    for row in week_table:
                        s = row["Sıra"]
                        if s <= 3:         rc = "#3498db"
                        elif s <= 6:       rc = "#f39c12"
                        elif s >= N - 2:   rc = "#e74c3c"
                        else:              rc = "#2d3748"

                        av = row["AV"]
                        av_color = "#27ae60" if av > 0 else ("#e74c3c" if av < 0 else "#a0aec0")

                        html_wt += (
                            f'<tr>'
                            f'<td><span style="background:{rc};color:#fff;border-radius:50%;'
                            f'display:inline-block;width:22px;height:22px;line-height:22px;'
                            f'font-size:11px;font-weight:700;text-align:center">{s}</span></td>'
                            f'<td class="left">{row["Takım"]}</td>'
                            f'<td>{row["O"]}</td><td>{row["G"]}</td>'
                            f'<td>{row["B"]}</td><td>{row["M"]}</td>'
                            f'<td>{row["AG"]}</td><td>{row["YG"]}</td>'
                            f'<td style="color:{av_color};font-weight:600">{av:+d}</td>'
                            f'<td style="font-weight:700">{row["Puan"]}</td>'
                            f'</tr>'
                        )
                    html_wt += "</tbody></table>"
                    st.markdown(html_wt, unsafe_allow_html=True)

                # Sağ: O haftanın maç sonuçları
                with right_col:
                    st.markdown(f"**{selected_week}. Hafta Maç Sonuçları**")
                    if not week_matches:
                        st.info("Bu haftada oynanan maç yok.")
                    else:
                        html_wm = '<table class="lig-table"><thead><tr>'
                        for h in ["Ev Sahibi","Skor","Deplasman"]:
                            cls = "left" if h in ["Ev Sahibi","Deplasman"] else ""
                            html_wm += f'<th class="{cls}">{h}</th>'
                        html_wm += "</tr></thead><tbody>"

                        for m in sorted(week_matches, key=lambda x: x["id"]):
                            hg, ag = m["hg"], m["ag"]
                            if hg > ag:
                                h_style = "font-weight:700;color:#27ae60"
                                a_style = "opacity:0.5"
                            elif hg < ag:
                                h_style = "opacity:0.5"
                                a_style = "font-weight:700;color:#27ae60"
                            else:
                                h_style = a_style = "color:#f39c12;font-weight:600"

                            html_wm += (
                                f'<tr>'
                                f'<td class="left" style="{h_style}">{m["home"]}</td>'
                                f'<td style="font-weight:700">{hg} - {ag}</td>'
                                f'<td class="left" style="{a_style}">{m["away"]}</td>'
                                f'</tr>'
                            )
                        html_wm += "</tbody></table>"
                        st.markdown(html_wm, unsafe_allow_html=True)

                    st.markdown("<br>**🔥 En Çok Gol Atılan İlk 10 Hafta**")

                    # Tüm oynanan haftaların gol toplamını hesapla
                    hafta_gol_map = {}
                    for m in played_fixtures:
                        w = m["week"]
                        hafta_gol_map[w] = hafta_gol_map.get(w, 0) + (m["hg"] or 0) + (m["ag"] or 0)

                    top10 = sorted(hafta_gol_map.items(), key=lambda x: -x[1])[:10]

                    html_top = '<table class="lig-table"><thead><tr>'
                    html_top += '<th>Sıra</th><th>Hafta</th><th>Gol</th><th>Maç Başı</th>'
                    html_top += "</tr></thead><tbody>"

                    for rank, (w, total_g) in enumerate(top10, 1):
                        mac_sayisi = len([m for m in played_fixtures if m["week"] == w])
                        mb = round(total_g / mac_sayisi, 2) if mac_sayisi else 0
                        # Seçili haftayı vurgula
                        row_style = "background:#2d3748;" if w == selected_week else ""
                        gol_color = "#27ae60" if rank == 1 else ("#f39c12" if rank <= 3 else "inherit")
                        html_top += (
                            f'<tr style="{row_style}">'
                            f'<td style="font-weight:700;color:#a0aec0">{rank}</td>'
                            f'<td>Hafta {w}</td>'
                            f'<td style="font-weight:700;color:{gol_color}">{total_g}</td>'
                            f'<td style="color:#718096">{mb}</td>'
                            f'</tr>'
                        )
                    html_top += "</tbody></table>"
                    st.markdown(html_top, unsafe_allow_html=True)


with tab6:
    st.markdown("### ⚙️ Ayarlar")
    st.markdown("#### 💾 Veritabanı & Dışa Aktarım")
    st.success(f"**SQLite:** `{DB_FILE}`", icon="🗄️")

    col_dl, col_xl, col_up = st.columns(3)
    with col_dl:
        if os.path.exists(DB_FILE):
            with open(DB_FILE, "rb") as f:
                st.download_button("⬇️ DB İndir (.db)", data=f.read(),
                                    file_name="lig.db", mime="application/octet-stream",
                                    use_container_width=True)
    with col_xl:
        if data["teams"]:
            try:
                excel_bytes = export_excel(data["teams"], data["fixtures"])
                st.download_button(
                    "📊 Excel İndir (.xlsx)",
                    data=excel_bytes,
                    file_name="lig_raporu.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True,
                    type="primary"
                )
            except Exception as e:
                st.error(f"Excel hatası: {e}")
        else:
            st.button("📊 Excel İndir", disabled=True, use_container_width=True,
                      help="Önce takım ekle")
    with col_up:
        uploaded = st.file_uploader("⬆️ JSON Yükle (eski format)", type="json",
                                     label_visibility="collapsed")
        if uploaded:
            try:
                loaded = json.load(uploaded)
                if "teams" in loaded and "fixtures" in loaded:
                    with get_db() as conn:
                        for t in loaded.get("teams", []):
                            conn.execute("INSERT OR IGNORE INTO teams (name) VALUES (?)", (t,))
                        for m in loaded.get("fixtures", []):
                            conn.execute(
                                "INSERT OR IGNORE INTO fixtures (id,week,devre,home,away,hg,ag,played) VALUES (?,?,?,?,?,?,?,?)",
                                (m["id"], m["week"], m.get("devre",1), m["home"], m["away"],
                                 m.get("hg"), m.get("ag"), 1 if m.get("played") else 0)
                            )
                    st.success("✅ JSON → DB yüklendi!")
                    st.rerun()
                else:
                    st.error("Geçersiz format.")
            except Exception as e:
                st.error(f"Hata: {e}")

    st.markdown("---")

    st.markdown("#### 🏷️ Takım Yönetimi")

    with st.expander("➕ Yeni Takım Ekle", expanded=len(data["teams"]) == 0):
        yeni_takim = st.text_input("Takım adı", key="yeni_takim", max_chars=40)
        if st.button("Ekle", key="takim_ekle_btn", type="primary"):
            yeni_takim = yeni_takim.strip()
            if not yeni_takim:
                st.error("Takım adı boş olamaz.")
            elif yeni_takim in data["teams"]:
                st.error("Bu takım zaten var.")
            else:
                db_save_team(yeni_takim)
                st.success(f"✅ {yeni_takim} eklendi!")
                st.rerun()

    st.markdown("<br>", unsafe_allow_html=True)

    if data["teams"]:
        st.markdown("**Mevcut Takımlar — İsim Düzenle**")
        st.caption("Değiştir butonuna bas → o takımın tüm maçları da güncellenir.")

        new_names = []
        cols_per_row = 3
        rows_needed = (len(data["teams"]) + cols_per_row - 1) // cols_per_row
        for row_i in range(rows_needed):
            cols = st.columns(cols_per_row)
            for col_i in range(cols_per_row):
                idx = row_i * cols_per_row + col_i
                if idx < len(data["teams"]):
                    val = cols[col_i].text_input(
                        f"{idx+1}.", value=data["teams"][idx],
                        key=f"tname_{idx}", max_chars=40
                    )
                    new_names.append((idx, val.strip()))

        if st.button("💾 İsimleri Kaydet", type="primary", use_container_width=True):
            changed_any = False
            for idx, new_name in new_names:
                if not new_name:
                    st.warning(f"{idx+1}. takım adı boş, atlandı.")
                    continue
                old_name = data["teams"][idx]
                if old_name != new_name:
                    db_rename_team(old_name, new_name)
                    changed_any = True
            if changed_any:
                st.success("✅ Güncellendi!")
                st.rerun()
            else:
                st.info("Değişiklik yok.")

        st.markdown("<br>", unsafe_allow_html=True)

        with st.expander("🗑️ Takım Sil"):
            st.warning("Takım silinince o takımın tüm maçları da silinir!", icon="⚠️")
            del_team = st.selectbox("Silinecek takım", data["teams"], key="del_team_sel")
            if st.button("🗑️ Takımı Sil", type="secondary", key="del_team_btn"):
                db_delete_team(del_team)
                st.success(f"{del_team} silindi.")
                st.rerun()

    st.markdown("---")

    st.markdown("#### 🗑️ Sıfırlama")
    col_r1, col_r2 = st.columns(2)
    with col_r1:
        if st.button("🗑️ Tüm Skorları Sıfırla", use_container_width=True):
            db_reset_scores()
            st.success("Skorlar sıfırlandı.")
            st.rerun()
    with col_r2:
        if st.button("💣 Her Şeyi Sıfırla", use_container_width=True, type="secondary"):
            if st.session_state.get("confirm_reset"):
                db_reset_all()
                st.session_state.confirm_reset = False
                st.success("Sıfırlandı.")
                st.rerun()
            else:
                st.session_state.confirm_reset = True
                st.warning("⚠️ Emin misin? Tekrar bas.")
